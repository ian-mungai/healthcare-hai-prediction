"""Lossless readers for bronze files that are not plain CSV: text lines, Excel sheet rows and SAS datasets.

Each reader streams one stored file into JSON Lines that Spark reads with an explicit schema, and returns counts only,
never values. Nothing is interpreted: text keeps every line with its terminator, Excel keeps each cell's stored value
as text with its cell type, and SAS keeps each variable's values as text. Splitting, typing and matching happen in
staging. PDFs keep every table row and text line pdfplumber reads, with the reader's version. CSV documents keep each
row's cells, Census-style JSON variable lists keep each variable's attributes, and Word documents keep each
paragraph and table row's text. Failure modes 68 to 80: data/lakehouse_planning/bronze_ipps_occmix_20261002/failure_modes.md;
97 to 110: publisher_dictionaries_20261002/; 113 to 129: bronze_county_20261002/; 132 to 134: bronze_remaining_20261003/.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
import zipfile
from collections.abc import Iterable, Iterator
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

# Lines end only at CRLF, LF or CR; str.splitlines would also split on form feeds and other separators [68].
LINE = re.compile(r"([^\r\n]*)(\r\n|\n|\r|)")
TEXT_COLUMNS = ("line_text", "line_terminator")
SHEET_SCHEMA = "r BIGINT, s STRING, i INT, w INT, d STRING, c ARRAY<STRING>, t ARRAY<STRING>"
SHEET_COLUMNS = (("sheet_name", "s"), ("sheet_index", "i"), ("sheet_row", "w"), ("date_system", "d"), ("cells", "c"), ("cell_types", "t"))
PDF_SCHEMA = "r BIGINT, k STRING, p INT, t INT, w INT, c ARRAY<STRING>, v STRING"
PDF_COLUMNS = (("row_kind", "k"), ("page", "p"), ("table_index", "t"), ("table_row", "w"), ("cells", "c"), ("reader", "v"))
CSV_ROWS_SCHEMA = "r BIGINT, c ARRAY<STRING>"
CSV_ROWS_COLUMNS = (("cells", "c"),)
JSON_SCHEMA = "r BIGINT, n STRING, a MAP<STRING,STRING>"
JSON_COLUMNS = (("variable_name", "n"), ("attributes", "a"))
DOCX_SCHEMA = "r BIGINT, k STRING, b INT, t INT, w INT, c ARRAY<STRING>"
DOCX_COLUMNS = (("block_kind", "k"), ("block_index", "b"), ("table_index", "t"), ("table_row", "w"), ("cells", "c"))
# Word's main part, its namespace and the largest uncompressed size read, so a crafted archive cannot exhaust memory [134].
DOCX_PART = "word/document.xml"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
DOCX_LIMIT = 200_000_000
BOM = b"\xef\xbb\xbf"
# xlrd cell type codes: empty and blank cells carry no value [73].
XLS_TYPES = {0: None, 1: "s", 2: "n", 3: "date", 4: "b", 5: "e", 6: None}


class ReaderError(ValueError):
    """A file cannot be stored losslessly; the message names the check, never a value."""


def split_lines(text: str) -> Iterator[tuple[str, str]]:
    """Yield each line and its terminator; a final terminator adds no empty line [68]."""
    position = 0
    while position < len(text):
        match = LINE.match(text, position)
        if match is None:  # pragma: no cover - the pattern always matches at least an empty line end
            raise ReaderError("the line pattern did not match")
        yield match.group(1), match.group(2)
        position = match.end()


def read_text_lines(path: Path, jsonl: Path, encodings: Iterable[str]) -> tuple[int, str]:
    """Store one row per line in the first allowed encoding that decodes strictly; the rows must rebuild the file."""
    data = path.read_bytes()
    if not data:
        raise ReaderError("the file is empty")
    expected = hashlib.sha256(data).hexdigest()
    failures = []
    for encoding in encodings:
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError as error:
            failures.append(f"{encoding.upper()} at byte {error.start}")
            continue
        rows = 0
        rebuilt = hashlib.sha256()
        with jsonl.open("w", encoding="ascii") as sink:
            for line, terminator in split_lines(text):
                rows += 1
                rebuilt.update((line + terminator).encode(encoding))
                sink.write(json.dumps({"r": rows, "v": [line, terminator]}, ensure_ascii=True) + "\n")
        if rebuilt.hexdigest() != expected:
            raise ReaderError("the stored lines do not rebuild the file byte for byte")
        return rows, encoding
    raise ReaderError(f"the file is not valid in any allowed encoding ({'; '.join(failures)})")


def cell_text(value: Any) -> str | None:
    """Return a cell's stored value as text without rounding or reformatting [72]."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return repr(value)
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    return str(value)


def trimmed(values: list[str | None], types: list[str | None]) -> tuple[list[str | None], list[str | None]]:
    """Drop trailing empty cells, which readers pad differently; inner gaps keep their positions."""
    while values and values[-1] is None:
        values.pop()
    return values, types[: len(values)]


def _xlsx_rows(path: Path) -> Iterator[tuple[str, int, int, list[str | None], list[str | None]]]:
    """Yield every row of every sheet with cached values, filling skipped rows so positions never shift [73] [74]."""
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        for index, sheet in enumerate(book.worksheets, 1):
            expected = 1
            for row in sheet.iter_rows(min_row=1):
                cells = [cell for cell in row if getattr(cell, "column", None) is not None and cell.value is not None]
                number = next((cell.row for cell in row if getattr(cell, "row", None) is not None), expected)
                while expected < number:
                    yield sheet.title, index, expected, [], []
                    expected += 1
                width = max((cell.column for cell in cells), default=0)
                values: list[str | None] = [None] * width
                types: list[str | None] = [None] * width
                for cell in cells:
                    values[cell.column - 1] = cell_text(cell.value)
                    types[cell.column - 1] = cell.data_type
                yield sheet.title, index, number, values, types
                expected = number + 1
    finally:
        book.close()


def _xlsx_formula_cells(path: Path) -> int:
    """Count formula cells, whose formula text bronze does not keep [74]."""
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=False)
    try:
        return sum(1 for sheet in book.worksheets for row in sheet.iter_rows() for cell in row if getattr(cell, "data_type", None) == "f")
    finally:
        book.close()


def _xls_rows(path: Path) -> tuple[Iterator[tuple[str, int, int, list[str | None], list[str | None]]], int]:
    """Return the rows of a legacy workbook and its date mode; dates stay serial numbers [75]."""
    import xlrd

    book = xlrd.open_workbook(str(path), on_demand=True)

    def rows() -> Iterator[tuple[str, int, int, list[str | None], list[str | None]]]:
        try:
            for index in range(book.nsheets):
                sheet = book.sheet_by_index(index)
                for number in range(sheet.nrows):
                    values: list[str | None] = []
                    types: list[str | None] = []
                    for code, value in zip(sheet.row_types(number), sheet.row_values(number), strict=True):
                        kind = XLS_TYPES[code]
                        types.append(kind)
                        if kind is None:
                            values.append(None)
                        elif kind == "b":
                            values.append("TRUE" if value else "FALSE")
                        elif kind == "e":
                            values.append(xlrd.error_text_from_code.get(value, str(value)))
                        else:
                            values.append(cell_text(value))
                    yield sheet.name, index + 1, number + 1, values, types
                book.unload_sheet(index)
        finally:
            book.release_resources()

    return rows(), int(book.datemode)


def read_sheet_rows(path: Path, jsonl: Path) -> tuple[int, str, dict[str, int]]:
    """Store one row per worksheet row with each cell's value as text and its type; return rows, label and counts."""
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        probe = load_workbook(path, read_only=True)
        system = "1904" if getattr(probe, "epoch", None) and probe.epoch.year == 1904 else "1900"
        probe.close()
        source: Iterable[tuple[str, int, int, list[str | None], list[str | None]]] = _xlsx_rows(path)
        stats = {"formula_cells": _xlsx_formula_cells(path)}
        # openpyxl never runs macros, and the macro project is not read [132].
        label = suffix[1:]
    elif suffix == ".xls":
        source, mode = _xls_rows(path)
        system = "1904" if mode else "1900"
        stats = {"formula_cells": 0}
        label = "xls"
    else:
        raise ReaderError(f"unsupported workbook extension {suffix!r}")
    rows = 0
    sheets = set()
    with jsonl.open("w", encoding="ascii") as sink:
        for name, index, number, values, types in source:
            values, types = trimmed(values, types)
            rows += 1
            sheets.add(index)
            record = {"r": rows, "s": name, "i": index, "w": number, "d": system, "c": values, "t": types}
            sink.write(json.dumps(record, ensure_ascii=True) + "\n")
    if not rows:
        raise ReaderError("the workbook has no rows")
    stats["sheets"] = len(sheets)
    return rows, label, stats


def sas_text(value: Any) -> str | None:
    """Return a SAS value as text: system missing is null, numbers keep their shortest round-trip form [77]."""
    if value is None:
        return None
    if isinstance(value, bytes):
        return value.decode("utf-8")
    if isinstance(value, str):
        return value
    number = float(value)
    if math.isnan(number):
        return None
    return repr(number)


def read_sas(path: Path, jsonl: Path) -> tuple[list[str], list[str | None], int, str]:
    """Store every SAS row with each variable's value as text; return names, labels, rows and encoding [77] [96]."""
    import pyreadstat

    data, meta = pyreadstat.read_sas7bdat(str(path), output_format="dict", disable_datetime_conversion=True, user_missing=True)
    headers = list(meta.column_names)
    labels = [label or None for label in (meta.column_labels or [None] * len(headers))]
    if len(labels) != len(headers):
        raise ReaderError("the SAS variable labels do not line up with the variable names")
    columns = [data[name] for name in headers]
    lengths = {len(column) for column in columns}
    if len(lengths) != 1:
        raise ReaderError("the SAS variables have different lengths")
    rows = lengths.pop()
    if meta.number_rows not in (None, rows) or meta.number_columns not in (None, len(headers)):
        raise ReaderError("the SAS row or column count differs from the file's metadata")
    with jsonl.open("w", encoding="ascii") as sink:
        for number in range(rows):
            sink.write(json.dumps({"r": number + 1, "v": [sas_text(column[number]) for column in columns]}, ensure_ascii=True) + "\n")
    return headers, labels, rows, str(meta.file_encoding or "")


def read_pdf_rows(path: Path, jsonl: Path) -> tuple[int, dict[str, int]]:
    """Store every table row and text line pdfplumber reads from a PDF, unchanged; refuse non-PDFs and text-free files [97] [98]."""
    with path.open("rb") as handle:
        if handle.read(5) != b"%PDF-":
            raise ReaderError("not a PDF")
    import pdfplumber

    reader = f"pdfplumber {pdfplumber.__version__}"
    rows = tables = lines = 0
    try:
        document = pdfplumber.open(path)
    except Exception as error:
        raise ReaderError(f"the PDF cannot be read ({type(error).__name__})") from None
    with document, jsonl.open("w", encoding="ascii") as sink:
        for page_number, page in enumerate(document.pages, 1):
            for table_number, table in enumerate(page.extract_tables(), 1):
                tables += 1
                for row_number, cells in enumerate(table, 1):
                    rows += 1
                    record = {
                        "r": rows,
                        "k": "table",
                        "p": page_number,
                        "t": table_number,
                        "w": row_number,
                        "c": [None if cell is None else str(cell) for cell in cells],
                        "v": reader,
                    }
                    sink.write(json.dumps(record, ensure_ascii=True) + "\n")
            for line_number, line in enumerate((page.extract_text() or "").splitlines(), 1):
                rows += 1
                lines += 1
                sink.write(
                    json.dumps({"r": rows, "k": "text", "p": page_number, "t": None, "w": line_number, "c": [line], "v": reader}, ensure_ascii=True) + "\n"
                )
        pages = len(document.pages)
    if not lines and not tables:
        raise ReaderError("the PDF has no text layer")
    return rows, {"pages": pages, "pdf_tables": tables, "text_lines": lines}


def split_quoted_line(text: str, delimiter: str) -> list[str] | None:
    """Split a line whose fields are all quoted or empty, where inner quotes are not doubled; None if it does not fit [130].

    A quote ends a field only when the delimiter or the line end follows it, so ``"Speak English "very well""`` keeps
    its inner quotes as published. Nothing is unescaped or dropped.
    """
    fields: list[str] = []
    index = 0
    while True:
        if index == len(text) or text[index] == delimiter:
            fields.append("")
            if index == len(text):
                return fields
            index += 1
            continue
        if text[index] != '"':
            return None
        search = index + 1
        while True:
            end = text.find('"', search)
            if end == -1:
                return None
            if end + 1 == len(text):
                fields.append(text[index + 1 : end])
                return fields
            if text[end + 1] == delimiter:
                fields.append(text[index + 1 : end])
                index = end + 2
                if index == len(text):
                    fields.append("")
                    return fields
                break
            search = end + 1


def parse_line(text: str, delimiter: str, quoted_fields: bool) -> list[str]:
    """Parse one physical line as strict CSV, or by the declared quoted-fields rule when strict CSV refuses it [130]."""
    try:
        rows = list(csv.reader([text], strict=True, delimiter=delimiter))
    except csv.Error as error:
        fields = split_quoted_line(text, delimiter) if quoted_fields else None
        if fields is None:
            raise ReaderError(f"the file is not valid CSV ({error})") from None
        return fields
    return rows[0] if rows else []


def read_csv_rows(
    path: Path, jsonl: Path, encodings: Iterable[str], delimiter: str = ",", byte_order_mark: bool = False, quoted_fields: bool = False
) -> tuple[int, str]:
    """Store each CSV row's cells as published, header included, with no width check [123] [130]."""
    data = path.read_bytes()
    if not data:
        raise ReaderError("the file is empty")
    if data.startswith(BOM):
        # A declared mark is removed and the rest must be UTF-8 [113].
        if not byte_order_mark:
            raise ReaderError("the file starts with a UTF-8 byte-order mark")
        encodings = ("utf-8-sig",)
    failures = []
    for encoding in encodings:
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError as error:
            failures.append(f"{encoding.upper()} at byte {error.start}")
            continue
        try:
            parsed = list(csv.reader(io.StringIO(text, newline=""), strict=True, delimiter=delimiter))
        except csv.Error as error:
            if not quoted_fields:
                raise ReaderError(f"the file is not valid CSV ({error})") from None
            # The declared rule reads the file line by line; lines strict CSV accepts are read strictly [130].
            parsed = [parse_line(line, delimiter, True) for line, _ in split_lines(text)]
        with jsonl.open("w", encoding="ascii") as sink:
            for number, row in enumerate(parsed, 1):
                sink.write(json.dumps({"r": number, "c": row}, ensure_ascii=True) + "\n")
        return len(parsed), encoding
    raise ReaderError(f"the file is not valid in any allowed encoding ({'; '.join(failures)})")


def read_json_variables(path: Path, jsonl: Path) -> int:
    """Store each variable of a Census-style variables object with its attributes as published text [123]."""
    try:
        document = json.loads(path.read_bytes().decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError):
        raise ReaderError("the file is not UTF-8 JSON") from None
    variables = document.get("variables") if isinstance(document, dict) else None
    if not isinstance(variables, dict) or not variables or not all(isinstance(value, dict) for value in variables.values()):
        raise ReaderError("the file has no Census-style variables object")
    rows = 0
    with jsonl.open("w", encoding="ascii") as sink:
        for name, attributes in variables.items():
            rows += 1
            # Non-text attributes (limits, nested lists) are kept as their JSON text, never dropped.
            values = {key: value if isinstance(value, str) else json.dumps(value, sort_keys=True) for key, value in attributes.items()}
            sink.write(json.dumps({"r": rows, "n": name, "a": values}, ensure_ascii=True) + "\n")
    return rows


def _paragraph_text(paragraph: ElementTree.Element) -> str:
    """Return a paragraph's visible text: run text, tabs and breaks; tab stops, deleted text and field codes are not text [133]."""
    parts: list[str] = []
    for run in paragraph.iter(f"{W}r"):
        for node in run:
            if node.tag == f"{W}t":
                parts.append(node.text or "")
            elif node.tag == f"{W}tab":
                parts.append("\t")
            elif node.tag in {f"{W}br", f"{W}cr"}:
                parts.append("\n")
    return "".join(parts)


def _docx_blocks(container: ElementTree.Element, counts: dict[str, int]) -> Iterator[tuple[str, int, int | None, int | None, list[str]]]:
    """Yield body paragraphs and table rows in document order with their block number; content controls are read through [133].

    A table is one block whose rows share its number. A cell's paragraphs, including those of a nested table, join with
    line breaks.
    """
    for child in container:
        if child.tag == f"{W}p":
            counts["blocks"] += 1
            yield "paragraph", counts["blocks"], None, None, [_paragraph_text(child)]
        elif child.tag == f"{W}tbl":
            counts["blocks"] += 1
            counts["tables"] += 1
            for number, row in enumerate(child.findall(f"{W}tr"), 1):
                cells = ["\n".join(_paragraph_text(paragraph) for paragraph in cell.iter(f"{W}p")) for cell in row.findall(f"{W}tc")]
                yield "table", counts["blocks"], counts["tables"], number, cells
        elif child.tag == f"{W}sdt":
            content = child.find(f"{W}sdtContent")
            if content is not None:
                yield from _docx_blocks(content, counts)


def read_docx_rows(path: Path, jsonl: Path) -> tuple[int, dict[str, int]]:
    """Store each body paragraph and table row of a Word document as text cells, in order; refuse anything else [133] [134]."""
    try:
        with zipfile.ZipFile(path) as archive:
            try:
                info = archive.getinfo(DOCX_PART)
            except KeyError:
                raise ReaderError(f"the package has no {DOCX_PART}") from None
            if info.file_size > DOCX_LIMIT:
                raise ReaderError(f"{DOCX_PART} is larger than {DOCX_LIMIT} bytes uncompressed")
            data = archive.read(info)
    except zipfile.BadZipFile:
        raise ReaderError("not a Word package (not a ZIP archive)") from None
    # Word never writes a DOCTYPE; refusing one means no entity can be declared, so the standard parser cannot expand
    # entities or fetch external ones, the attacks defusedxml guards against.
    if not data.removeprefix(BOM).startswith(b"<"):
        raise ReaderError(f"{DOCX_PART} is not UTF-8 XML")
    if b"<!DOCTYPE" in data:
        raise ReaderError(f"{DOCX_PART} declares a DOCTYPE, which Word never writes")
    parser: ElementTree.XMLPullParser = ElementTree.XMLPullParser(events=("end",))
    try:
        parser.feed(data)
        parser.close()
    except ElementTree.ParseError:
        raise ReaderError(f"{DOCX_PART} is not well-formed XML") from None
    # The root element closes last.
    root: ElementTree.Element | None = None
    for event in parser.read_events():
        element = event[-1]
        if isinstance(element, ElementTree.Element):
            root = element
    body = root.find(f"{W}body") if root is not None and root.tag == f"{W}document" else None
    if body is None:
        raise ReaderError(f"{DOCX_PART} has no document body")
    rows = paragraphs = 0
    counts = {"blocks": 0, "tables": 0}
    with jsonl.open("w", encoding="ascii") as sink:
        for kind, block, table, row, cells in _docx_blocks(body, counts):
            rows += 1
            paragraphs += kind == "paragraph"
            sink.write(json.dumps({"r": rows, "k": kind, "b": block, "t": table, "w": row, "c": cells}, ensure_ascii=True) + "\n")
    if not rows:
        raise ReaderError("the document has no paragraphs or tables")
    return rows, {"paragraphs": paragraphs, "docx_tables": counts["tables"]}
