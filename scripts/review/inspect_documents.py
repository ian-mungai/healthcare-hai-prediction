"""Inspect Office and PDF structures with bundled readers, emitting no cell/text examples."""

from __future__ import annotations

import argparse
import json
import logging
import re
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

import openpyxl
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from scripts.review.review_remaining import digest, safe_path, write_json

log = logging.getLogger(__name__)


def workbook(path: Path) -> dict[str, Any]:
    """Check every worksheet's declared dimensions and a bounded cell prefix."""
    book = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        sheets = []
        for sheet in book.worksheets:
            counts: Counter[str] = Counter()
            for row in sheet.iter_rows(max_row=min(sheet.max_row or 1000, 1000), max_col=min(sheet.max_column or 1000, 1000)):
                for cell in row:
                    if cell.value is None:
                        counts["blank"] += 1
                    elif cell.data_type == "f":
                        counts["formula"] += 1
                    elif cell.data_type == "e":
                        counts["excel_error"] += 1
                    elif cell.data_type == "n":
                        counts["numeric"] += 1
                    else:
                        counts["other"] += 1
            sheets.append({"sheet": sheet.title, "declared_rows": sheet.max_row, "declared_columns": sheet.max_column, "sample_cell_types": dict(counts)})
        return {
            "status": "workbook_structure_sampled",
            "sheets": sheets,
            "limits": "At most 1000 rows and columns per sheet; dimensions may be formatting. Formula results and table boundaries not validated.",
        }
    finally:
        book.close()


def pdf(path: Path) -> dict[str, Any]:
    """Extract every page locally and report only safe definition-term counts."""
    reader = PdfReader(path, strict=False)
    if reader.is_encrypted:
        return {"status": "pending_encrypted_pdf"}
    terms: Counter[str] = Counter()
    empty_pages = 0
    for page in reader.pages:
        text = page.extract_text() or ""
        empty_pages += not text.strip()
        for term in ("suppressed", "confidence interval", "per 100,000", "per 1,000", "standardized infection ratio", "fiscal year", "calendar year"):
            terms[term] += len(re.findall(re.escape(term), text, flags=re.IGNORECASE))
    return {
        "status": "pdf_text_structure_checked",
        "pages": len(reader.pages),
        "pages_without_extracted_text": empty_pages,
        "definition_term_counts": dict(terms),
        "limits": "Extraction only; no visual/table/semantic validation. No dates or units attributed to measures merely because a term occurs.",
    }


def inspect(profile_root: Path, source_root: Path, output: Path) -> dict[str, Any]:
    """Recheck immutable artifacts before applying the document readers."""
    results = []
    cache: dict[str, dict[str, Any]] = {}
    for source_file in sorted(profile_root.glob("*.json")):
        source = json.loads(source_file.read_text())
        for candidate in source["candidates"]:
            for artifact in candidate.get("artifacts", []):
                suffix = Path(artifact["path"]).suffix.lower()
                if suffix not in {".pdf", ".xlsx", ".xlsm", ".docx"} or artifact["integrity"] != "verified":
                    continue
                path = safe_path(source_root, artifact["path"])
                result: dict[str, Any]
                if digest(path) != artifact["sha256"]:
                    result = {"status": "artifact_drifted"}
                elif artifact["sha256"] in cache:
                    result = cache[artifact["sha256"]]
                else:
                    try:
                        if suffix in {".xlsx", ".xlsm"}:
                            result = workbook(path)
                        elif suffix == ".pdf":
                            result = pdf(path)
                        else:
                            with zipfile.ZipFile(path) as archive:
                                xml = archive.read("word/document.xml")
                            result = {"status": "docx_structure_checked", "table_nodes": xml.count(b"<w:tbl>"), "limits": "No semantic/visual review."}
                    except (OSError, ValueError, KeyError, zipfile.BadZipFile, PdfReadError) as error:
                        result = {"status": "document_parse_failed", "error_type": type(error).__name__}
                    cache[artifact["sha256"]] = result
                results.append({"source_id": source["source_id"], "path": artifact["path"], "sha256": artifact["sha256"], "result": result})
        log.info("Document structure checked: %s", source["source_id"])
    report = {
        "code_sha256": digest(Path(__file__)),
        "dependencies": {"openpyxl": openpyxl.__version__},
        "results": results,
        "statuses": dict(Counter(r["result"]["status"] for r in results)),
    }
    write_json(output, report)
    return report


def main() -> None:
    """Run document structure checks from the bundled Python runtime."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--profiles", type=Path, required=True)
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    logging.basicConfig(level=logging.ERROR)
    inspect(args.profiles, args.source_root, args.output)


if __name__ == "__main__":
    main()
